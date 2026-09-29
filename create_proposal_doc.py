import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                      f'<w:top w:w="{top}" w:type="dxa"/>'
                      f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
                      f'<w:left w:w="{left}" w:type="dxa"/>'
                      f'<w:right w:w="{right}" w:type="dxa"/>'
                      f'</w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def format_run(run, font_name="TH Sarabun New", size_pt=16, bold=False, color_rgb=(0, 0, 0)):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.font.color.rgb = RGBColor(*color_rgb)
    rPr = run._r.get_or_add_rPr()
    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{font_name}" w:hAnsi="{font_name}" w:cs="{font_name}"/>')
    rPr.append(rFonts)

def add_styled_paragraph(doc, text="", font_name="TH Sarabun New", size_pt=16, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=0):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if text:
        run = p.add_run(text)
        format_run(run, font_name=font_name, size_pt=size_pt, bold=bold)
    return p

def create_table_for_model(doc, model_title, columns, data_rows):
    p_title = add_styled_paragraph(doc, model_title, bold=True, space_before=10, space_after=4)
    
    table = doc.add_table(rows=len(data_rows) + 1, cols=len(columns))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    
    col_widths = [Inches(1.8), Inches(2.2), Inches(1.8), Inches(1.2)]
    
    # Header row
    hdr_cells = table.rows[0].cells
    for i, col_name in enumerate(columns):
        hdr_cells[i].text = col_name
        set_cell_background(hdr_cells[i], "E5E7EB")  # light gray
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=150, right=150)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        if len(p.runs) > 0:
            format_run(p.runs[0], bold=True, size_pt=15)
            
    # Data rows
    for r_idx, row_data in enumerate(data_rows):
        row_cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = str(val)
            set_cell_margins(row_cells[c_idx], top=100, bottom=100, left=150, right=150)
            p = row_cells[c_idx].paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            if c_idx in [0, 1]:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if len(p.runs) > 0:
                format_run(p.runs[0], bold=False, size_pt=14)
                
    # Set widths
    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width
            
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def build_proposal_document(file_path):
    doc = Document()
    
    # Page setup (A4 margins: 1 inch / 2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        
    # Title
    add_styled_paragraph(doc, "ข้อเสนอเค้าโครงโครงงาน", size_pt=20, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_styled_paragraph(doc, "(Class Project Proposal)", size_pt=17, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18)
    
    # ผู้พัฒนา
    add_styled_paragraph(doc, "ผู้พัฒนา", size_pt=16, bold=True, space_after=2)
    add_styled_paragraph(doc, "1. xxxxxxxxxx (ใส่ชื่อ-นามสกุล และรหัสนักศึกษา)", size_pt=16, space_after=14)
    
    # ชื่อโปรแกรม
    p_prog = doc.add_paragraph()
    p_prog.paragraph_format.space_after = Pt(14)
    p_prog.paragraph_format.line_spacing = 1.15
    run1 = p_prog.add_run("ชื่อโปรแกรม  ")
    format_run(run1, size_pt=16, bold=True)
    run2 = p_prog.add_run("ส่งต่อของดีในสาขา (ICT UBU Share)")
    format_run(run2, size_pt=16, bold=False)
    
    # รายละเอียดฟังก์ชัน/ความสามารถของโปรแกรม
    add_styled_paragraph(doc, "รายละเอียดฟังก์ชัน/ความสามารถของโปรแกรม", size_pt=16, bold=True, space_after=6)
    
    features = [
        "1. ระบบจัดการบัญชีผู้ใช้งาน (Authentication & Profile Management): รองรับการสมัครสมาชิก, เข้าสู่ระบบ, ออกจากระบบ, เปลี่ยนและรีเซ็ตรหัสผ่าน พร้อมระบบจัดการโปรไฟล์นักศึกษา (รหัส นศ., ชั้นปี, รูปโปรไฟล์, แนะนำตัว และช่องทางติดต่อ Line/เบอร์โทร)",
        "2. ระบบลงประกาศสิ่งของส่งต่อ (Item Management): สมาชิกสามารถลงประกาศสิ่งของที่ต้องการส่งต่อหรือแจกฟรี (ระบุชื่อของ, รายละเอียดสภาพ, ราคา หรือตั้ง 0 เพื่อแจกฟรี, ช่องทางติดต่อ และรูปภาพประกอบ) พร้อมฟังก์ชันแก้ไขและลบประกาศของตนเอง",
        "3. ระบบสืบค้นและคัดกรองสิ่งของ (Search & Filter): ผู้ใช้สามารถค้นหาสิ่งของตามชื่อ/รายละเอียด และกรองดูตามสถานะ (พร้อมส่งต่อ, มีคนขอรับแล้ว, ส่งมอบสำเร็จ) ได้อย่างสะดวกรวดเร็ว",
        "4. ระบบตรวจจับช่องทางติดต่ออัตโนมัติ (Smart Contact Links): ระบบแปลงลิงก์ช่องทางติดต่ออัตโนมัติ (เบอร์โทร, LINE, Facebook, IG, X, อีเมล) เป็นปุ่มคลิกติดต่อเจ้าของได้ทันที",
        "5. ระบบการขอนัดรับสิ่งของ (Meetup & Claim Request): ผู้สนใจสามารถส่งคำขอนัดรับของ โดยเลือกจุดนัดรับใต้ตึกสาขา/คณะ (เช่น ใต้ตึกคณะ ICT ลานกิจกรรม, หน้าแล็บ CS ชั้น 2, โรงอาหารกลาง ฯลฯ) พร้อมระบุวันที่นัดรับและข้อความถึงผู้ให้ โดยระบบจะปรับสถานะของเป็น 'มีคนขอรับแล้ว' อัตโนมัติ",
        "6. ระบบจัดการคำขอของฉันและยืนยันการส่งมอบ (My Requests & Handover Confirmation): มีหน้าติดตามรายการของตนเองและคำขอนัดรับ โดยเจ้าของของสามารถกดยืนยันการส่งมอบสำเร็จเพื่อเปลี่ยนสถานะเป็น 'ส่งมอบสำเร็จ' ได้",
        "7. ระบบความคิดเห็นและการตอบกลับ (Item Comments & Nested Replies): ผู้ใช้งานสามารถแสดงความคิดเห็นใต้ประกาศเพื่อสอบถามรายละเอียดเพิ่มเติม และรองรับการตอบกลับความคิดเห็น (Nested Replies)",
        "8. ระบบการแจ้งเตือนแบบเรียลไทม์ (Notification System): แจ้งเตือนเมื่อมีคำขอนัดรับใหม่, เมื่อเจ้าของยืนยันการส่งมอบ, หรือเมื่อมีผู้มาคอมเมนต์/ตอบกลับ พร้อมไอคอนกระดิ่งและสถานะการอ่าน",
        "9. ระบบแดชบอร์ดสรุปสถิติภาพรวม (Admin Analytics Dashboard): แสดงสถิติจำนวนสิ่งของทั้งหมด, ของที่พร้อมส่งต่อ, ของที่มีคนขอรับ, ของที่ส่งมอบสำเร็จ และสัดส่วนของแจกฟรีเทียบกับของมีราคา สำหรับผู้ดูแลระบบ"
    ]
    
    for feat in features:
        p_f = doc.add_paragraph()
        p_f.paragraph_format.left_indent = Inches(0.2)
        p_f.paragraph_format.space_after = Pt(3)
        p_f.paragraph_format.line_spacing = 1.15
        run = p_f.add_run(feat)
        format_run(run, size_pt=15)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    
    # โมเดลข้อมูล/ตารางฐานข้อมูล
    add_styled_paragraph(doc, "โมเดลข้อมูล/ตารางฐานข้อมูล", size_pt=17, bold=True, space_after=8)
    
    cols = ["ชื่อฟิลด์", "ความหมาย", "ชนิด/ความยาว", "หมายเหตุ"]
    
    # Model 1: Item
    item_rows = [
        ["id", "รหัสสิ่งของ", "AutoField (PK)", "อัตโนมัติ"],
        ["owner (FK:User)", "ผู้ลงประกาศสิ่งของ", "ForeignKey (User)", "เชื่อมกับตาราง User"],
        ["title", "ชื่อสิ่งของ", "CharField (200)", ""],
        ["price", "ราคา (0 = แจกฟรี)", "DecimalField (7, 2)", "ค่าเริ่มต้น 0"],
        ["contact_link", "ลิงก์ช่องทางติดต่อ", "CharField (300)", "ว่างได้"],
        ["image", "รูปภาพประกอบสิ่งของ", "ImageField", "ว่างได้"],
        ["description", "รายละเอียดสภาพของ/โน้ต", "TextField", ""],
        ["status", "สถานะของสิ่งของ", "CharField (20)", "ค่าเริ่มต้น 'available'"],
        ["created_at", "วันเวลาที่ลงประกาศ", "DateTimeField", "อัตโนมัติ (auto_now_add)"],
    ]
    create_table_for_model(doc, "โมเดล/ตาราง 1: Item (สิ่งของสำหรับส่งต่อ)", cols, item_rows)
    
    # Model 2: ClaimRequest
    claim_rows = [
        ["id", "รหัสคำขอนัดรับ", "AutoField (PK)", "อัตโนมัติ"],
        ["item (FK:Item)", "สิ่งของที่ขอรับ", "ForeignKey (Item)", "เชื่อมกับตาราง Item"],
        ["requester (FK:User)", "ผู้ขอรับสิ่งของ", "ForeignKey (User)", "เชื่อมกับตาราง User"],
        ["pickup_place", "จุดนัดรับใต้ตึกสาขา", "CharField (150)", "ตัวเลือกจุดนัดรับในคณะ"],
        ["meetup_date", "วันที่นัดรับของ", "DateField", ""],
        ["note", "ข้อความถึงเจ้าของ/ผู้ให้", "TextField", "ว่างได้"],
        ["is_confirmed", "เจ้าของยืนยันส่งมอบสำเร็จ", "BooleanField", "ค่าเริ่มต้น False"],
        ["created_at", "วันเวลาที่ส่งคำขอ", "DateTimeField", "อัตโนมัติ (auto_now_add)"],
    ]
    create_table_for_model(doc, "โมเดล/ตาราง 2: ClaimRequest (การขอนัดรับของ)", cols, claim_rows)
    
    # Model 3: Comment
    comment_rows = [
        ["id", "รหัสความคิดเห็น", "AutoField (PK)", "อัตโนมัติ"],
        ["item (FK:Item)", "สิ่งของที่แสดงความคิดเห็น", "ForeignKey (Item)", "เชื่อมกับตาราง Item"],
        ["author (FK:User)", "ผู้แสดงความคิดเห็น", "ForeignKey (User)", "เชื่อมกับตาราง User"],
        ["parent (FK:Comment)", "ความคิดเห็นหลักที่ตอบกลับ", "ForeignKey (Self)", "ว่างได้ (Nested Reply)"],
        ["content", "ข้อความความคิดเห็น", "TextField", ""],
        ["created_at", "วันเวลาที่แสดงความคิดเห็น", "DateTimeField", "อัตโนมัติ (auto_now_add)"],
    ]
    create_table_for_model(doc, "โมเดล/ตาราง 3: Comment (ความคิดเห็น/ถาม-ตอบ)", cols, comment_rows)
    
    # Model 4: Notification
    notif_rows = [
        ["id", "รหัสการแจ้งเตือน", "AutoField (PK)", "อัตโนมัติ"],
        ["recipient (FK:User)", "ผู้รับการแจ้งเตือน", "ForeignKey (User)", "เชื่อมกับตาราง User"],
        ["sender (FK:User)", "ผู้ส่งการแจ้งเตือน", "ForeignKey (User)", "ว่างได้"],
        ["item (FK:Item)", "สิ่งของที่เกี่ยวข้อง", "ForeignKey (Item)", "ว่างได้"],
        ["claim (FK:ClaimRequest)", "คำขอนัดรับที่เกี่ยวข้อง", "ForeignKey (ClaimRequest)", "ว่างได้"],
        ["comment (FK:Comment)", "ความคิดเห็นที่เกี่ยวข้อง", "ForeignKey (Comment)", "ว่างได้"],
        ["notification_type", "ประเภทการแจ้งเตือน", "CharField (50)", "claim_new, reply_new ฯลฯ"],
        ["title", "หัวข้อการแจ้งเตือน", "CharField (200)", ""],
        ["message", "ข้อความแจ้งเตือน", "TextField", ""],
        ["is_read", "สถานะการอ่าน", "BooleanField", "ค่าเริ่มต้น False"],
        ["created_at", "วันเวลาที่แจ้งเตือน", "DateTimeField", "อัตโนมัติ (auto_now_add)"],
    ]
    create_table_for_model(doc, "โมเดล/ตาราง 4: Notification (การแจ้งเตือนในระบบ)", cols, notif_rows)
    
    # Model 5: UserProfile
    profile_rows = [
        ["id", "รหัสโปรไฟล์", "AutoField (PK)", "อัตโนมัติ"],
        ["user (OneToOne:User)", "บัญชีผู้ใช้งานระบบ", "OneToOneField (User)", "เชื่อม 1:1 กับ User"],
        ["student_id", "รหัสนักศึกษา", "CharField (20)", "ว่างได้"],
        ["year_level", "ชั้นปี (1 - 4)", "IntegerField", "ค่าเริ่มต้น 2"],
        ["phone_number", "เบอร์โทรศัพท์", "CharField (20)", "ว่างได้"],
        ["line_id", "LINE ID", "CharField (50)", "ว่างได้"],
        ["bio", "คำแนะนำตัวสั้นๆ", "TextField", "ว่างได้"],
        ["avatar", "รูปโปรไฟล์", "ImageField", "ว่างได้"],
        ["theme_mode", "โหมดธีม (Light/Dark)", "CharField (20)", "ค่าเริ่มต้น 'light'"],
        ["accent_color", "สีธีมหลักของระบบ", "CharField (20)", "ค่าเริ่มต้น 'blue'"],
        ["notify_email", "รับแจ้งเตือนทางอีเมล", "BooleanField", "ค่าเริ่มต้น True"],
        ["show_contact_publicly", "แสดงข้อมูลติดต่อสาธารณะ", "BooleanField", "ค่าเริ่มต้น True"],
    ]
    create_table_for_model(doc, "โมเดล/ตาราง 5: UserProfile (ข้อมูลโปรไฟล์นักศึกษาและการตั้งค่า)", cols, profile_rows)
    
    # Model 6: User (auth_user)
    user_rows = [
        ["id", "รหัสผู้ใช้งาน", "AutoField (PK)", "อัตโนมัติ"],
        ["username", "ชื่อผู้ใช้สำหรับล็อกอิน", "CharField (150)", "ไม่ซ้ำกัน (Unique)"],
        ["password", "รหัสผ่านที่เข้ารหัสแล้ว", "CharField (128)", "Password Hash"],
        ["first_name", "ชื่อจริง", "CharField (150)", "ว่างได้"],
        ["last_name", "นามสกุล", "CharField (150)", "ว่างได้"],
        ["email", "ที่อยู่อีเมล", "EmailField (254)", ""],
        ["is_staff", "สิทธิ์เจ้าหน้าที่/ผู้ดูแล", "BooleanField", "เข้าดูหน้า Dashboard สถิติ"],
        ["is_active", "สถานะเปิดใช้งานบัญชี", "BooleanField", "ค่าเริ่มต้น True"],
        ["date_joined", "วันเวลาที่สมัครสมาชิก", "DateTimeField", "อัตโนมัติ (auto_now_add)"],
    ]
    create_table_for_model(doc, "โมเดล/ตาราง 6: User (auth_user - บัญชีผู้ใช้งานระบบ)", cols, user_rows)
    
    doc.save(file_path)
    print("Successfully generated proposal document.")

if __name__ == "__main__":
    out_dir = r"c:\Users\Asus By Kristana\Desktop\Project_wab"
    target_path = os.path.join(out_dir, "ข้อเสนอเค้าโครงโครงงาน_ส่งต่อของดีในสาขา.docx")
    build_proposal_document(target_path)
